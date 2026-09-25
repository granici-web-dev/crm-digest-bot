from datetime import date, datetime
from io import BytesIO
from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from syrupy.assertion import SnapshotAssertion

from digest.config import AppConfig
from digest.metrics.frame import prepare_lead_frame
from digest.reports.context import ReportContext
from digest.reports.modules import IMPLEMENTED_MODULES
from digest.reports.render import ReportLanguage
from factories import BUCHAREST, make_snapshot_row

MONTH_END = date(2026, 9, 30)
MONTHLY_MODULES = ("m2", "m3", "m4", "m5", "m8", "m19")
PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
# Синтетические данные клиента: кадр их не несёт, тест подкладывает их, чтобы проверить, что
# ни текст, ни вложение не берут лишних колонок.
CLIENT_DATA = {
    "name": "CLIENT_TEST_NUME",
    "phone": "+40700000001",
    "email": "client1@example.com",
    "notes": "NOTA_CLIENT_TEST",
}


def at(day: date, hour: int) -> datetime:
    return datetime(day.year, day.month, day.day, hour, tzinfo=BUCHAREST)


def lead(lead_id: int, created_at: datetime, **overrides: Any) -> dict[str, Any]:
    return make_snapshot_row(
        **{"lead_id": lead_id, "created_at": created_at, "last_contact_at": created_at, **overrides}
    )


def month_frame(app_config: AppConfig) -> pd.DataFrame:
    september, august = date(2026, 9, 10), date(2026, 8, 12)
    rows = [
        *(
            lead(lead_id, at(september, 12), showroom="Brașov", assigned_to_id=8)
            for lead_id in range(1, 9)
        ),
        lead(9, at(september, 13), showroom="Brașov", assigned_to_id=9, ofertat=True),
        lead(
            10,
            at(september, 14),
            showroom="Brașov",
            assigned_to_id=9,
            category="WON",
            status_name="Clienți",
            ofertat=True,
            converted_at=at(date(2026, 9, 20), 12),
        ),
        lead(11, at(september, 15), showroom="Cluj", assigned_to_id=10, ofertat=True),
        lead(
            12,
            at(september, 16),
            showroom="Cluj",
            assigned_to_id=10,
            category="LOST",
            loss_reason="BUGET",
            status_name="BUGET",
            status_changed_at=at(date(2026, 9, 18), 12),
        ),
        lead(
            13,
            at(september, 17),
            showroom=None,
            assigned_to_id=12,
            category="LOST",
            loss_reason="IRELEVANT",
            status_name="IRELEVANT",
        ),
        lead(
            14,
            at(august, 12),
            showroom="București",
            assigned_to_id=12,
            category="LOST",
            loss_reason="BUGET",
            status_name="BUGET",
            status_changed_at=at(date(2026, 8, 20), 12),
        ),
        lead(
            15,
            at(august, 12),
            showroom="București",
            assigned_to_id=13,
            category="LOST",
            loss_reason="NU_RASPUNS",
            status_name="NU A RASPUNS",
            status_changed_at=at(date(2026, 8, 25), 12),
        ),
        lead(
            16,
            at(date(2026, 5, 5), 12),
            showroom="București",
            assigned_to_id=13,
            category="WON",
            status_name="Clienți",
            converted_at=at(date(2026, 7, 7), 12),
        ),
    ]
    return prepare_lead_frame(rows, app_config).assign(**CLIENT_DATA)


def context(app_config: AppConfig, language: ReportLanguage) -> ReportContext:
    return ReportContext(MONTH_END, None, None, app_config, "sofabelle", language)


@pytest.mark.parametrize("language", ["ro", "ru"])
@pytest.mark.parametrize("module_id", MONTHLY_MODULES)
def test_monthly_module_text(
    app_config: AppConfig, language: ReportLanguage, module_id: str, snapshot: SnapshotAssertion
) -> None:
    result = IMPLEMENTED_MODULES[module_id](month_frame(app_config), context(app_config, language))

    assert result.text == snapshot
    for value in CLIENT_DATA.values():
        assert value not in result.text


def module_text(app_config: AppConfig, module_id: str) -> str:
    return IMPLEMENTED_MODULES[module_id](month_frame(app_config), context(app_config, "ro")).text


def test_funnel_text_and_photo(app_config: AppConfig) -> None:
    result = IMPLEMENTED_MODULES["m2"](month_frame(app_config), context(app_config, "ro"))

    assert "Brașov: 10 → 10 → 2 → 1" in result.text
    assert "Cluj: 2 → 2 → 1 → 0" in result.text
    assert "1 lead-uri fără showroom, dintre care 1 irelevante" in result.text
    assert "(fără showroom)" not in result.text
    assert result.alerts == ()
    assert "București" not in result.text
    assert "Total: 13 → 12 → 3 → 1" in result.text
    assert "Lead-uri → utile → oferte → Clienți" in result.text
    assert "Clienți: din lead-urile lunii, status la data raportului." in result.text
    assert "Contracte" not in result.text
    assert result.photo is not None
    assert result.photo.filename == "funnel_2026-09.png"
    assert result.photo.content.startswith(PNG_SIGNATURE)


def test_useful_lead_without_showroom_alerts_ops_without_client_data(
    app_config: AppConfig,
) -> None:
    frame = pd.concat(
        [
            month_frame(app_config),
            prepare_lead_frame(
                [lead(17, at(date(2026, 9, 21), 12), showroom=None)], app_config
            ).assign(**CLIENT_DATA),
        ]
    )

    result = IMPLEMENTED_MODULES["m2"](frame, context(app_config, "ro"))

    assert "2 lead-uri fără showroom, dintre care 1 irelevante" in result.text
    assert result.alerts == (
        "m2 за 09.2026: 1 полезных лидов без шоурума (всего без шоурума 2). "
        "Заполнить поле Showroom в mefi.",
    )


def test_trend_text_and_photo(app_config: AppConfig) -> None:
    result = IMPLEMENTED_MODULES["m3"](month_frame(app_config), context(app_config, "ro"))

    assert "Lead-uri: Apr 0 · Mai 1 · Iun 0 · Iul 0 · Aug 2 · Sep 13" in result.text
    assert "Contracte: Apr 0 · Mai 0 · Iun 0 · Iul 1 · Aug 0 · Sep 1" in result.text
    assert result.photo is not None
    assert result.photo.filename == "trend_2026-09.png"
    assert result.photo.content.startswith(PNG_SIGNATURE)


def test_scr_text_shows_level_and_target(app_config: AppConfig) -> None:
    text = module_text(app_config, "m4")

    assert "țintă ≥10,0%" in text
    assert "Brașov 10,0% · Elită ✓" in text
    assert "Cluj 0,0% · Sub minim ✗" in text
    assert "(fără showroom) —\n" in text
    assert "Total 8,3% · Bine ✗" in text


def test_cockpit_text_uses_code_lines_without_pre_or_spi(app_config: AppConfig) -> None:
    text = module_text(app_config, "m5")

    assert "<pre>" not in text
    assert "SPI" not in text
    assert "<b>Moaca Andreea</b> · Brașov" in text
    assert "<code>Lead-uri 2 · Utile 2 · Oferte 2 · Clienți 1</code>" in text
    assert "<i>Clienți: din lead-urile lunii, status la data raportului</i>" in text
    assert "Contracte" not in text
    assert "<code>SCR 50,0% ✓ · L2O 100,0% ✓ · O2C 50,0% ✓" in text
    assert text.count("<code>") == 6 * 3


def test_loss_reasons_text_compares_with_previous_month(app_config: AppConfig) -> None:
    text = module_text(app_config, "m8")

    assert "Motive de pierdere: Sep față de Aug" in text
    assert "Pierdute 2 (Aug 2) (=)" in text
    assert "Buget 1 · 50% (Aug 1) (=)" in text
    assert "Irelevant 1 · 50% (Aug 0) (—)" in text
    assert "Nu a răspuns 0 · 0% (Aug 1) (−100%)" in text


def workbook(app_config: AppConfig) -> tuple[str, Workbook]:
    result = IMPLEMENTED_MODULES["m19"](month_frame(app_config), context(app_config, "ro"))
    assert result.document is not None
    return result.document.filename, load_workbook(BytesIO(result.document.content))


def sheet_rows(book: Workbook, name: str) -> list[tuple[Any, ...]]:
    return list(book[name].iter_rows(values_only=True))


def row_labelled(rows: list[tuple[Any, ...]], label: str) -> tuple[Any, ...]:
    return next(row for row in rows if row[0] == label)


CLIENTI_NOTE = "Clienți: din lead-urile lunii, status la data raportului"


def test_monthly_workbook_sheets_and_filename(app_config: AppConfig) -> None:
    filename, book = workbook(app_config)

    assert filename == "sofabelle_2026-09.xlsx"
    assert book.sheetnames == [
        "KPI consilieri",
        "Funnel showroom",
        "Motive pierdere",
        "Lead-uri luna",
    ]


def test_monthly_workbook_manager_sheet(app_config: AppConfig) -> None:
    _, book = workbook(app_config)
    rows = sheet_rows(book, "KPI consilieri")

    assert rows[0][:7] == (
        "Consilier",
        "Showroom",
        "Lead-uri",
        "Utile",
        "Oferte",
        "Clienți",
        "SCR",
    )
    assert rows[1][0] == "Țintă"
    assert rows[1][6] == "≥10,0%"
    moaca = next(row for row in rows if row[0] == "Moaca Andreea")
    assert moaca[1:6] == ("Brașov", 2, 2, 2, 1)
    assert moaca[6] == pytest.approx(0.5)
    godja = next(row for row in rows if row[0] == "Godja Adina Maria")
    assert godja[6] == pytest.approx(0.0)
    raileanu = next(row for row in rows if row[0] == "Raileanu  Leon")
    assert raileanu[2] == 0
    assert raileanu[6] == "—"
    assert len(rows) == 2 + 6 + 2
    assert rows[-1][0] == CLIENTI_NOTE


def test_monthly_workbook_funnel_and_loss_sheets(app_config: AppConfig) -> None:
    _, book = workbook(app_config)

    funnel = sheet_rows(book, "Funnel showroom")
    assert funnel[0] == (
        "Showroom",
        "Lead-uri",
        "Utile",
        "Oferte",
        "Clienți",
        "SCR",
        "L2O",
        "O2C",
    )
    assert funnel[1][:5] == ("Brașov", 10, 10, 2, 1)
    total = row_labelled(funnel, "Total")
    assert total[:5] == ("Total", 13, 12, 3, 1)
    assert total[5] == pytest.approx(1 / 12)
    assert funnel[-1][0] == CLIENTI_NOTE

    losses = sheet_rows(book, "Motive pierdere")
    assert losses[0] == ("Motiv", "Sep", "Pondere", "Aug", "Variație")
    assert losses[-1][:4] == ("Total", 2, 1.0, 2)
    assert losses[-1][4] == pytest.approx(0.0)
    irelevant = next(row for row in losses if row[0] == "Irelevant")
    assert irelevant[1:] == (1, 0.5, 0, "—")


def test_monthly_workbook_is_romanian_in_russian_run(app_config: AppConfig) -> None:
    result = IMPLEMENTED_MODULES["m19"](month_frame(app_config), context(app_config, "ru"))
    assert result.document is not None
    losses = sheet_rows(load_workbook(BytesIO(result.document.content)), "Motive pierdere")

    assert losses[0] == ("Motiv", "Sep", "Pondere", "Aug", "Variație")
    assert {row[0] for row in losses[1:]} == {"Buget", "Irelevant", "Nu a răspuns", "Total"}


def test_monthly_workbook_lead_sheet_matches_funnel_without_client_data(
    app_config: AppConfig,
) -> None:
    _, book = workbook(app_config)
    rows = sheet_rows(book, "Lead-uri luna")

    assert rows[0][:8] == (
        "ID",
        "Creat",
        "Zi",
        "Showroom",
        "Sursa",
        "Status",
        "Ofertat",
        "Consilier",
    )
    assert rows[0][9] == "Luna: lead-uri create 31.08.2026 19:00 → 30.09.2026 19:00"
    assert len(rows) - 1 == 13
    cells = {
        str(cell)
        for sheet in book.worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
    }
    for value in CLIENT_DATA.values():
        assert value not in cells
