from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from digest.config import AppConfig
from digest.metrics.frame import prepare_lead_frame
from digest.reports.context import ReportContext
from digest.reports.modules.weekly_excel import excel_attachment_report
from factories import BUCHAREST, make_snapshot_row

SUNDAY = date(2026, 9, 27)
MONDAY = date(2026, 9, 21)
PERSONAL_DATA_HEADERS = {"name", "phone", "email", "nume", "telefon", "e-mail", "client"}


def at(day: date, hour: int) -> datetime:
    return datetime(day.year, day.month, day.day, hour, tzinfo=BUCHAREST)


def lead(lead_id: int, created_at: datetime, **overrides: Any) -> dict[str, Any]:
    return make_snapshot_row(**{"lead_id": lead_id, "created_at": created_at, **overrides})


def week_frame(app_config: AppConfig) -> pd.DataFrame:
    rows = [
        lead(1, at(MONDAY, 11), showroom="Brașov", source_name="Site", ofertat=True),
        lead(2, at(MONDAY, 12), showroom="Brașov", source_name="WhatsApp"),
        lead(3, at(SUNDAY, 12), showroom="Cluj", source_name="Site"),
        lead(4, at(SUNDAY, 13), showroom=None, source_name=None, assigned_to_id=99),
        lead(5, at(MONDAY, 14), showroom="Cluj", source_name="Showroom"),
        lead(6, at(SUNDAY + timedelta(days=1), 11)),
    ]
    return prepare_lead_frame(rows, app_config)


def workbook(app_config: AppConfig, tenant_id: str = "sofabelle") -> tuple[str, Workbook]:
    context = ReportContext(SUNDAY, None, None, app_config, tenant_id, "ro")
    result = excel_attachment_report(week_frame(app_config), context)
    assert result.document is not None
    return result.document.filename, load_workbook(BytesIO(result.document.content))


def sheet_rows(book: Workbook, name: str) -> list[tuple[Any, ...]]:
    return list(book[name].iter_rows(values_only=True))


def test_workbook_has_manual_report_sheets_plus_leads_and_visits(app_config: AppConfig) -> None:
    filename, book = workbook(app_config)

    assert filename == "sofabelle_sapt39_2026.xlsx"
    assert book.sheetnames == [
        "Свод день-шоурум",
        "По дням шоурум-источник",
        "Шоурум-источник итог",
        "Lead-uri",
        "Vizite",
    ]


def test_filename_takes_tenant_id(app_config: AppConfig) -> None:
    filename, _ = workbook(app_config, tenant_id="alt_tenant")

    assert filename == "alt_tenant_sapt39_2026.xlsx"


def test_summary_sheets_have_manual_headers_and_totals(app_config: AppConfig) -> None:
    _, book = workbook(app_config)

    summary = sheet_rows(book, "Свод день-шоурум")
    assert summary[2] == (
        "Zi lucrătoare",
        "Brașov",
        "București",
        "Cluj",
        "(fără showroom)",
        "TOTAL",
    )
    assert summary[-1] == ("TOTAL", 2, 0, 1, 1, 4)

    by_source = sheet_rows(book, "Шоурум-источник итог")
    assert by_source[2] == ("Showroom", "Site", "WhatsApp", "(fără sursă)", "TOTAL")
    assert by_source[-1] == ("TOTAL", 2, 1, 1, 4)

    detail = sheet_rows(book, "По дням шоурум-источник")
    day_totals = [row[-1] for row in detail if str(row[0]).startswith("Total zi")]
    assert sum(day_totals) == 4


def test_lead_sheets_match_summary_and_carry_no_personal_data(app_config: AppConfig) -> None:
    _, book = workbook(app_config)

    leads = sheet_rows(book, "Lead-uri")
    visits = sheet_rows(book, "Vizite")

    header = ("ID", "Creat", "Zi lucrătoare", "Showroom", "Sursa", "Status", "Ofertat", "Consilier")
    assert leads[0] == header
    assert visits[0] == ("ID", "Creat", "Zi", *header[3:])
    assert not {cell.lower() for cell in header} & PERSONAL_DATA_HEADERS
    assert [row[0] for row in leads[1:]] == [1, 2, 3, 4]
    assert leads[1][1:3] == (datetime(2026, 9, 21, 11, 0), datetime(2026, 9, 21, 0, 0))
    assert leads[1][6:] == ("✅DA", "Dragoi Mihaela")
    assert leads[4][3:5] == (None, None)
    assert leads[4][7] == "id 99"
    assert [row[0] for row in visits[1:]] == [5]
