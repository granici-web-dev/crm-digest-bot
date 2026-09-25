"""Build tests/fixtures/etalon-2026-05.json from docs/reference/SB KPi.xlsx.

Usage: uv run python scripts/build_etalon.py

The workbook supplies leads and thresholds only. expected_by_agent is computed here by
the brief formulas (ADR-002, docs/kpi-definitions.md), independently of metrics/;
the workbook's own 03_KPI_Agenti values go to excel_reference for comparison.

Prints only counts; lead data never goes to stdout (the workbook holds client contacts).
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
import warnings
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "reference" / "SB KPi.xlsx"
MANAGERS = ROOT / "config" / "managers.yaml"
TARGET = ROOT / "tests" / "fixtures" / "etalon-2026-05.json"
TZ = ZoneInfo("Europe/Bucharest")

LEADS_SHEET = "01_Input_Leads"
SETTINGS_SHEET = "02_Setari_Targete"
KPI_SHEET = "03_KPI_Agenti"
KPI_HEADER_ROW = 3
KPI_ROWS = range(4, 10)
KPI_FIRST_COL, KPI_LAST_COL = 3, 31  # C..AE

# Excel header -> etalon field. Nume, Telefon, E-mail, Informatii, Revenire 1-3, Oraș,
# Interacțiuni and helper columns V..AG are deliberately absent.
LEAD_FIELDS = {
    "Data ultimei solicitări": "created_at",
    "Ultimul contact": "last_contact_at",
    "Status": "status",
    "Ofertat": "ofertat",
    "Desemnat": "assigned_to",
    "Showroom": "showroom",
    "Sursa": "source",
    "Data revenire": "data_revenire",
    "UTM_Source": "utm_source",
    "UTM_Campanie": "utm_campaign",
    "UTM_Content": "utm_content",
    "UTM_Medium": "utm_medium",
}
DATETIME_FIELDS = {"created_at", "last_contact_at"}
DATE_FIELDS = {"data_revenire"}
PII_HEADERS = ("Nume", "Telefon", "E-mail")
# Single-word client names coincide with consultant first names and source labels, so
# substring search is limited to identifying shapes: full names, e-mails, phone digits.
PII_SUBSTRING_MIN_LEN = 6
PHONE_MIN_DIGITS = 7

# Deliberately duplicated from config/status-mapping.yaml: the etalon is an oracle for
# metrics/, so it must not share category code or config parsing with it (ADR-002).
STATUS_CLIENTI = "Clienți"
STATUSES_IRELEVANT = {"IRELEVANT", "SPAM"}
STATUSES_PARTNERSHIP = {"DESIGNER", "INFLUENCER"}
STATUS_NU_A_RASPUNS = "NU A RASPUNS"
STATUS_BUGET = "BUGET"
STATUS_PRODUS_NEPOTRIVIT = "PRODUS NEPOTRIVIT"
SOURCE_SHOWROOM = "Showroom"
OFERTAT_YES = "✅DA"
ACTIVE_OFFER_STALE_DAYS = 14
# Doja Ovidiu has leads in the workbook but no mefi user (docs/PLAN.md, «Ждём извне»).
SYNTHETIC_MANAGER_IDS = {"Doja Ovidiu": 999}

# Counted by hand from the fixture leads on 2026-09-25 (ADR-002). The build fails if the
# computed values drift from this record.
MANUAL_CHECK: dict[str, Any] = {
    "agent": "Moaca Andreea",
    "checked_on": "2026-09-25",
    "note": "23 строки в Excel, 1 INFLUENCER исключён; клиентов нет, пересечение SC не проверено",
    "counts": {
        "leads": 22,
        "irr_leads": 3,
        "useful": 19,
        "clienti": 0,
        "offers": 7,
        "nar": 5,
        "buget": 4,
        "pnp": 2,
        "showroom_visits": 10,
        "clienti_from_showroom": 0,
        "active_offers_14": 6,
    },
}


def snake(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", ascii_text.lower()).strip("_")


def parse_datetime(value: Any) -> str | None:
    if value in (None, ""):
        return None
    local = datetime.strptime(str(value), "%d-%m-%Y %H:%M:%S").replace(tzinfo=TZ)
    return local.isoformat()


def parse_date(value: Any) -> str | None:
    if value in (None, ""):
        return None
    return datetime.strptime(str(value)[:10], "%d-%m-%Y").date().isoformat()


def read_leads(ws: Any) -> tuple[list[dict[str, Any]], set[str]]:
    header = [cell.value for cell in ws[1]]
    col = {name: i for i, name in enumerate(header) if name is not None}
    missing = [h for h in (*LEAD_FIELDS, *PII_HEADERS) if h not in col]
    if missing:
        raise SystemExit(f"{LEADS_SHEET}: missing columns {missing}")

    leads: list[dict[str, Any]] = []
    pii: set[str] = set()
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        for h in PII_HEADERS:
            if row[col[h]] not in (None, ""):
                pii.add(str(row[col[h]]).strip())
        if row[col["Status"]] in (None, ""):
            continue
        lead: dict[str, Any] = {"row": row_no}
        for header_name, field in LEAD_FIELDS.items():
            value = row[col[header_name]]
            if field in DATETIME_FIELDS:
                value = parse_datetime(value)
            elif field in DATE_FIELDS:
                value = parse_date(value)
            elif value == "":
                value = None
            lead[field] = value
        leads.append(lead)
    return leads, pii


def read_thresholds(ws: Any) -> tuple[str, dict[str, float]]:
    analysis_date = ws["B3"].value
    if not isinstance(analysis_date, datetime):
        raise SystemExit(f"{SETTINGS_SHEET}!B3 is not a date")
    thresholds: dict[str, float] = {}
    for row in range(4, ws.max_row + 1):
        label, value = ws.cell(row, 1).value, ws.cell(row, 2).value
        if label is None or value is None:
            continue
        thresholds[snake(str(label))] = value
    return analysis_date.date().isoformat(), thresholds


def read_excel_reference(ws: Any) -> dict[str, dict[str, Any]]:
    headers = {
        c: snake(str(ws.cell(KPI_HEADER_ROW, c).value))
        for c in range(KPI_FIRST_COL, KPI_LAST_COL + 1)
    }
    expected: dict[str, dict[str, Any]] = {}
    for row in KPI_ROWS:
        agent = ws.cell(row, 1).value
        values: dict[str, Any] = {"showroom": ws.cell(row, 2).value}
        values.update({key: ws.cell(row, c).value for c, key in headers.items()})
        expected[str(agent)] = values
    return expected


def collapse_spaces(name: str) -> str:
    return " ".join(name.split())


def manager_ids_by_name(agents: list[str]) -> dict[str, int]:
    roster = yaml.safe_load(MANAGERS.read_text(encoding="utf-8"))["managers"]
    by_name = {collapse_spaces(m["name"]): m["id"] for m in roster}
    ids = {agent: by_name.get(agent, SYNTHETIC_MANAGER_IDS.get(agent)) for agent in agents}
    unknown = [agent for agent, manager_id in ids.items() if manager_id is None]
    if unknown:
        raise SystemExit(f"no id in managers.yaml for {unknown}")
    return {agent: manager_id for agent, manager_id in ids.items() if manager_id is not None}


def ratio(numerator: int, denominator: int) -> float | None:
    return None if denominator == 0 else numerator / denominator


def is_stale_offer(lead: dict[str, Any], analysis_date: date) -> bool:
    if lead["last_contact_at"] is None:
        return False
    contact_day = datetime.fromisoformat(lead["last_contact_at"]).date()
    return (analysis_date - contact_day).days > ACTIVE_OFFER_STALE_DAYS


def brief_counts(leads: list[dict[str, Any]], analysis_date: date) -> dict[str, int]:
    counts = dict.fromkeys(MANUAL_CHECK["counts"], 0)
    for lead in leads:
        status = lead["status"]
        if lead["created_at"] is None or status in STATUSES_PARTNERSHIP:
            continue
        clienti = status == STATUS_CLIENTI
        irelevant = status in STATUSES_IRELEVANT
        offer = lead["ofertat"] == OFERTAT_YES
        showroom = lead["source"] == SOURCE_SHOWROOM
        counts["leads"] += 1
        counts["irr_leads"] += irelevant
        counts["clienti"] += clienti
        counts["offers"] += offer
        counts["nar"] += status == STATUS_NU_A_RASPUNS
        counts["buget"] += status == STATUS_BUGET
        counts["pnp"] += status == STATUS_PRODUS_NEPOTRIVIT
        counts["showroom_visits"] += showroom
        counts["clienti_from_showroom"] += clienti and showroom
        counts["active_offers_14"] += (
            offer and not clienti and not irelevant and is_stale_offer(lead, analysis_date)
        )
    counts["useful"] = counts["leads"] - counts["irr_leads"]
    return counts


def brief_kpis(c: dict[str, int]) -> dict[str, float | None]:
    return {
        "scr": ratio(c["clienti"], c["useful"]),
        "l2o": ratio(c["offers"], c["useful"]),
        "o2c": ratio(c["clienti"], c["offers"]),
        "cdr": ratio(c["leads"] - c["nar"], c["leads"]),
        "plr": ratio(c["buget"], c["leads"] - c["irr_leads"] - c["nar"]),
        "sc": ratio(c["clienti_from_showroom"], c["showroom_visits"]),
        "pfr": ratio(c["pnp"], c["useful"]),
        "acr": ratio(c["active_offers_14"], c["leads"]),
        "irr": ratio(c["irr_leads"], c["leads"]),
    }


def expected_by_agent(
    leads: list[dict[str, Any]], agents: list[str], analysis_date: date
) -> dict[str, dict[str, Any]]:
    expected: dict[str, dict[str, Any]] = {}
    for agent in agents:
        # 03_KPI_Agenti keys agents by TRIM(Desemnat); mefi keeps "Raileanu  Leon" with two spaces.
        own = [lead for lead in leads if collapse_spaces(lead["assigned_to"] or "") == agent]
        counts = brief_counts(own, analysis_date)
        expected[agent] = {"counts": counts, "kpi": brief_kpis(counts)}
    return expected


def leaked_pii(payload: Any, pii: set[str], dump: str) -> list[str]:
    leaks: list[str] = []

    def walk(node: Any) -> None:
        if isinstance(node, dict):
            for v in node.values():
                walk(v)
        elif isinstance(node, list):
            for v in node:
                walk(v)
        elif node is not None and str(node).strip() in pii:
            leaks.append("exact")

    walk(payload)
    dump_digits = re.sub(r"\D", "", dump)
    for p in pii:
        digits = re.sub(r"\D", "", p)
        if len(digits) >= PHONE_MIN_DIGITS and digits in dump_digits:
            leaks.append("phone")
        elif len(p) >= PII_SUBSTRING_MIN_LEN and (" " in p or "@" in p) and p in dump:
            leaks.append("substring")
    return leaks


def main() -> None:
    # openpyxl warns about unsupported Excel extensions (conditional formatting); irrelevant here.
    warnings.filterwarnings("ignore", module="openpyxl")
    wb = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)

    leads, pii = read_leads(wb[LEADS_SHEET])
    analysis_date, thresholds = read_thresholds(wb[SETTINGS_SHEET])
    excel_reference = read_excel_reference(wb[KPI_SHEET])
    expected = expected_by_agent(leads, list(excel_reference), date.fromisoformat(analysis_date))
    checked = expected[MANUAL_CHECK["agent"]]["counts"]
    if checked != MANUAL_CHECK["counts"]:
        sys.exit(f"manual check drifted for {MANUAL_CHECK['agent']}, etalon not written")
    created = sorted(lead["created_at"] for lead in leads if lead["created_at"])

    payload = {
        "meta": {
            "source_file": SOURCE.name,
            "analysis_date": analysis_date,
            "timezone": TZ.key,
            "date_range": {"from": created[0], "to": created[-1]},
            "lead_count": len(leads),
            "formulas": "docs/brief.md §3, docs/kpi-definitions.md (ADR-002)",
            "manual_check": MANUAL_CHECK,
        },
        "thresholds": thresholds,
        "leads": leads,
        # Agent name as in 03_KPI_Agenti (spaces collapsed) -> mefi assigned_to.id.
        "managers": manager_ids_by_name(list(excel_reference)),
        "expected_by_agent": expected,
        # SB KPi.xlsx 03_KPI_Agenti as-is, including the AC defect (ACR = 0 everywhere).
        # Comparison only; tests assert against expected_by_agent.
        "excel_reference": excel_reference,
    }
    dump = json.dumps(payload, ensure_ascii=False, indent=2)

    leaks = leaked_pii(payload, pii, dump)
    if leaks:
        # Count only: printing the matched value would itself leak it.
        sys.exit(f"PII check failed: {len(leaks)} match(es), etalon not written")

    TARGET.write_text(dump + "\n", encoding="utf-8")
    print(
        f"leads={len(leads)} thresholds={len(thresholds)} agents={len(expected)} "
        f"pii_values_checked={len(pii)} -> {TARGET.relative_to(ROOT)}"
    )


if __name__ == "__main__":
    main()
